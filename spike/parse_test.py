"""Phase 0: native-Python parsing of PPTX / DOCX / PDF into normalised units, plus the vision heuristic
and slide composites (no PowerPoint / LibreOffice).

Usage:
    python spike/parse_test.py                 # generate synthetic samples and parse them
    python spike/parse_test.py file1.pptx ...  # parse real files
"""
import hashlib
import io
import json
import sys
from collections import Counter
from pathlib import Path

import pymupdf as fitz
from docx import Document as DocxDocument
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt

from omml import omml_to_latex

HERE = Path(__file__).parent
SAMPLES = HERE / "samples"
OUT = HERE / "out"

# Vision heuristic thresholds (Settings in the real app).
MIN_IMAGES = 2
BIG_IMAGE_AREA = 0.40
LOW_TEXT_WORDS = 15
MIN_IMAGE_AREA = 0.02       # smaller than this = icon/logo, ignored
TEMPLATE_IMAGE_RATIO = 0.5  # image on >50% of slides = template decoration
VECTOR_DIAGRAM_PATHS = 40   # PDF pages with more vector paths than this likely hold a drawn diagram


# ---------------------------------------------------------------- samples

def _chart(label, color):
    img = Image.new("RGB", (640, 420), "white")
    d = ImageDraw.Draw(img)
    d.line([(50, 370), (610, 370)], fill="black", width=3)
    d.line([(50, 370), (50, 30)], fill="black", width=3)
    d.line([(50, 370), (580, 60)], fill=color, width=5)
    d.text((240, 385), "time (s)", fill="black")
    d.text((60, 35), label, fill="black")
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def _logo():
    img = Image.new("RGB", (80, 80), "navy")
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def make_samples():
    SAMPLES.mkdir(exist_ok=True)
    logo = _logo()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    def slide(title, body=None, notes=None):
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = title
        s.shapes.add_picture(io.BytesIO(logo), Inches(12.4), Inches(0.1), Inches(0.7), Inches(0.7))
        if body:
            tb = s.shapes.add_textbox(Inches(0.7), Inches(1.6), Inches(6), Inches(4))
            tb.text_frame.word_wrap = True
            tb.text_frame.text = body
            for p in tb.text_frame.paragraphs:
                p.font.size = Pt(20)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    slide("Kinematics: definitions",
          "Displacement is a vector: change in position. Velocity is the rate of change of displacement. "
          "Acceleration is the rate of change of velocity. SI units: m, m/s, m/s^2.",
          notes="Emphasise the vector nature of displacement versus distance.")
    s = slide("Velocity-time vs displacement-time graphs", "Compare the gradients.")
    s.shapes.add_picture(io.BytesIO(_chart("v (m/s)", "red")), Inches(0.7), Inches(2.6), Inches(5.8), Inches(3.8))
    s.shapes.add_picture(io.BytesIO(_chart("s (m)", "blue")), Inches(6.8), Inches(2.6), Inches(5.8), Inches(3.8))
    s = slide("Projectile motion")
    s.shapes.add_picture(io.BytesIO(_chart("y (m)", "green")), Inches(1.5), Inches(1.3), Inches(10), Inches(5.8))
    t = slide("SUVAT equations").shapes.add_table(3, 2, Inches(0.7), Inches(1.6), Inches(8), Inches(2)).table
    for r, (a, b) in enumerate([("Equation", "Missing"), ("v = u + at", "s"), ("s = ut + 0.5at^2", "v")]):
        t.cell(r, 0).text, t.cell(r, 1).text = a, b
    prs.save(SAMPLES / "kinematics.pptx")

    doc = DocxDocument()
    doc.add_heading("Ohm's Law", 1)
    doc.add_paragraph("Ohm's law states V = IR, where V is voltage, I is current and R is resistance.")
    doc.add_heading("Series circuits", 2)
    doc.add_paragraph("In series, resistances add: R_total = R1 + R2 + ...")
    doc.add_picture(io.BytesIO(_chart("V (volts)", "purple")), width=Inches(4))
    tbl = doc.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text, tbl.cell(0, 1).text, tbl.cell(1, 0).text, tbl.cell(1, 1).text = "R1", "10 ohm", "R2", "20 ohm"
    doc.save(SAMPLES / "circuits.docx")

    pdf = fitz.open()
    p = pdf.new_page()
    p.insert_text((72, 72), "Chapter 2: Motion in One Dimension", fontsize=16)
    p.insert_text((72, 100), "Average velocity is displacement divided by elapsed time.", fontsize=11)
    p = pdf.new_page()
    p.insert_text((72, 72), "Figure 2.1", fontsize=11)
    p.insert_image(fitz.Rect(72, 90, 520, 400), stream=_chart("x (m)", "orange"))
    pdf.save(SAMPLES / "textbook_ch2.pdf")
    return [SAMPLES / "kinematics.pptx", SAMPLES / "circuits.docx", SAMPLES / "textbook_ch2.pdf"]


# ---------------------------------------------------------------- XML text (shared by PPTX/DOCX)

M_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"


def _local(tag):
    return tag.rsplit("}", 1)[-1] if isinstance(tag, str) else ""


def xml_text(el):
    """Text of an element in document order. Equations become $LaTeX$; mc:Fallback (equation images) is skipped."""
    out = []

    def walk(e):
        name = _local(e.tag)
        if name == "Fallback":
            return
        if e.tag == M_NS + "oMath":
            out.append(" $" + omml_to_latex(e) + "$ ")
            return
        if name == "t" and e.text and not e.tag.startswith(M_NS):
            out.append(e.text)
        elif name in ("br", "tab"):
            out.append(" ")
        for c in e:
            walk(c)

    walk(el)
    return "".join(out).strip()


def count_equations(el):
    return len(el.findall(".//" + M_NS + "oMath"))


# ---------------------------------------------------------------- PPTX

def _iter_shapes(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(sh.shapes)
        else:
            yield sh


def _font(size):
    try:
        return ImageFont.truetype("arial.ttf", size)
    except OSError:
        return ImageFont.load_default()


def compose_slide(prs, slide, out_png, width=1600):
    """Approximate slide render: pictures at true position/size, text boxes drawn in place."""
    scale = width / prs.slide_width
    canvas = Image.new("RGB", (width, int(prs.slide_height * scale)), "white")
    draw = ImageDraw.Draw(canvas)
    for sh in _iter_shapes(slide.shapes):
        if sh.left is None or sh.width is None:
            continue
        box = (int(sh.left * scale), int(sh.top * scale), int(sh.width * scale), int(sh.height * scale))
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            try:
                im = Image.open(io.BytesIO(sh.image.blob)).convert("RGB")
                canvas.paste(im.resize((max(box[2], 1), max(box[3], 1))), box[:2])
            except Exception:  # EMF/WMF etc. that Pillow can't open: draw a placeholder
                draw.rectangle([box[0], box[1], box[0] + box[2], box[1] + box[3]], outline="gray")
        elif sh.has_text_frame and sh.text_frame.text.strip():
            draw.multiline_text((box[0] + 4, box[1] + 4), sh.text_frame.text, fill="black", font=_font(24))
    canvas.save(out_png)


def parse_pptx(path):
    prs = Presentation(path)
    area = prs.slide_width * prs.slide_height
    slides = []
    for i, s in enumerate(prs.slides, 1):
        texts, pics = [], []
        # Walk the slide XML (not python-pptx shapes) so equations and AlternateContent text are included.
        tree = s.shapes._spTree
        tables = list(tree.iter("{*}tbl"))
        for tbl in tables:
            rows = [" | ".join(xml_text(c) for c in r.iter("{*}tc")) for r in tbl.iter("{*}tr")]
            texts.append("[table]\n" + "\n".join(rows))
        for para in tree.iter("{*}p"):
            if _local(para.tag) != "p":
                continue
            if any(a in tables or _local(a.tag) == "Fallback" for a in para.iterancestors()):
                continue
            t = xml_text(para)
            if t:
                texts.append(t)
        for sh in _iter_shapes(s.shapes):
            # Filled picture placeholders are <p:pic> too, but report shape_type PLACEHOLDER.
            if _local(sh._element.tag) == "pic" and hasattr(sh, "image"):
                blob = sh.image.blob
                pics.append({"sha": hashlib.sha1(blob).hexdigest()[:12], "ext": sh.image.ext,
                             "area": (sh.width * sh.height) / area, "blob": blob})
        notes = s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide else ""
        slides.append({"index": i, "text": "\n".join(texts), "notes": notes, "pics": pics, "equations": count_equations(tree)})

    # Images repeated on most slides are template decoration (logos, footers).
    seen = Counter(p["sha"] for sl in slides for p in {q["sha"]: q for q in sl["pics"]}.values())
    template = {h for h, n in seen.items() if len(slides) > 2 and n / len(slides) > TEMPLATE_IMAGE_RATIO}

    units = []
    stem = Path(path).stem
    for sl in slides:
        content = [p for p in sl["pics"] if p["sha"] not in template and p["area"] >= MIN_IMAGE_AREA]
        words = len(sl["text"].split())
        reasons = []
        if len(content) >= MIN_IMAGES:
            reasons.append(f"{len(content)} images")
        if any(p["area"] > BIG_IMAGE_AREA for p in content):
            reasons.append("large image")
        if content and words < LOW_TEXT_WORDS:
            reasons.append(f"low text ({words} words)")
        unit = {"type": "slide", "index": sl["index"], "text": sl["text"], "notes": sl["notes"],
                "images": len(content), "ignored_images": len(sl["pics"]) - len(content), "equations": sl["equations"],
                "needs_vision": bool(reasons), "reasons": reasons}
        if reasons:
            png = OUT / f"{stem}_slide{sl['index']}.png"
            compose_slide(prs, prs.slides[sl["index"] - 1], png)
            unit["composite"] = png.name
        units.append(unit)
    return units


# ---------------------------------------------------------------- DOCX

def parse_docx(path):
    doc = DocxDocument(path)
    units, cur = [], {"type": "section", "heading": "(intro)", "text": [], "images": 0}
    body = doc.element.body
    for el in body.iterchildren():
        tag = el.tag.split("}")[1]
        if tag == "p":
            style = el.find(".//{*}pStyle")
            style = style.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val") if style is not None else ""
            text = xml_text(el)
            imgs = len(el.findall(".//{*}blip"))
            if style.lower().startswith("heading") and text:
                if cur["text"] or cur["images"]:
                    units.append(cur)
                cur = {"type": "section", "heading": text, "level": style, "text": [], "images": 0}
            elif text:
                cur["text"].append(text)
            cur["images"] += imgs
        elif tag == "tbl":
            rows = [" | ".join(xml_text(c) for c in r.iter("{*}tc"))
                    for r in el.iter("{*}tr")]
            cur["text"].append("[table]\n" + "\n".join(rows))
    units.append(cur)
    for u in units:
        u["text"] = "\n".join(u["text"])
        u["needs_vision"] = u["images"] > 0
    return units


# ---------------------------------------------------------------- PDF

def parse_pdf(path):
    pdf = fitz.open(path)
    units, stem = [], Path(path).stem
    for p in pdf:
        text = p.get_text("text").strip()
        page_area = p.rect.width * p.rect.height
        words = len(text.split())
        rects = [r for info in p.get_images(full=True) for r in p.get_image_rects(info[0])]
        full = [r for r in rects if r.width * r.height > 0.9 * page_area]
        # A full-page image under real text is a background/template (or a scan with an OCR layer).
        img_area = sum(r.width * r.height for r in rects if r not in full or words < 20)
        drawings = len(p.get_drawings())
        reasons = []
        if words < 20 and full:
            reasons.append("scanned page")
        elif not text:
            reasons.append("no text layer")
        if img_area / page_area > 0.15 and not (words < 20 and full):
            reasons.append(f"figure {img_area / page_area:.0%} of page")
        if drawings > VECTOR_DIAGRAM_PATHS:
            reasons.append(f"vector diagram ({drawings} paths)")
        unit = {"type": "page", "index": p.number + 1, "text": text, "needs_vision": bool(reasons), "reasons": reasons}
        if reasons:
            png = OUT / f"{stem}_p{p.number + 1}.png"
            p.get_pixmap(dpi=110).save(png)
            unit["render"] = png.name
        units.append(unit)
    return units


PARSERS = {".pptx": parse_pptx, ".docx": parse_docx, ".pdf": parse_pdf}


def main():
    OUT.mkdir(exist_ok=True)
    files = [Path(a) for a in sys.argv[1:]] or make_samples()
    for f in files:
        units = PARSERS[f.suffix.lower()](f)
        (OUT / f"{f.stem}.units.json").write_text(json.dumps(units, indent=2), encoding="utf-8")
        print(f"\n== {f.name}: {len(units)} units, {sum(u['needs_vision'] for u in units)} need vision")
        for u in units:
            label = u.get("heading") or f"{u['type']} {u.get('index')}"
            flag = f"VISION ({', '.join(u.get('reasons', [])) or 'embedded images'})" if u["needs_vision"] else "text"
            print(f"  - {label}: {len(u['text'].split())} words, {flag}")


if __name__ == "__main__":
    main()
