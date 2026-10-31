"""Layouts for showing an original document in the app (not for the model).

PDFs are rendered client-side by pdf.js straight from the stored file. Slides and Word documents have no native
renderer here, so they are described as positioned/ordered blocks of real text plus pictures; the UI lays them out as
HTML so the text stays selectable. Pictures are written once to assets/<doc id>/view/ (browser-safe formats only)."""
import io
from pathlib import Path

from docx import Document as DocxDocument
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn

from . import config
from .parsers import _local, docx_blocks, docx_groups, sha, xml_text

EMU_PT = 12700
WEB_IMAGE = {"png", "jpg", "jpeg", "gif", "webp", "bmp", "svg"}
W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _picture(did: int, blob: bytes, ext: str) -> str | None:
    """Asset path for a picture the browser can show; other formats (EMF, WMF, TIFF) are converted to PNG."""
    ext = ext.lower().lstrip(".")
    h = sha(blob)[:20]
    out_dir = config.assets_dir() / str(did) / "view"
    for e in (ext, "png"):
        if (out_dir / f"{h}.{e}").exists():
            return f"{did}/view/{h}.{e}"
    out_dir.mkdir(parents=True, exist_ok=True)
    if ext in WEB_IMAGE:
        (out_dir / f"{h}.{ext}").write_bytes(blob)
        return f"{did}/view/{h}.{ext}"
    try:
        im = Image.open(io.BytesIO(blob))
        im.load()
        im.save(out_dir / f"{h}.png", "PNG")
        return f"{did}/view/{h}.png"
    except Exception:
        return None


# ---------------------------------------------------------------- PPTX

def _identity(x, y, w, h):
    return x, y, w, h


def _group_transform(grp, outer):
    """Children of a group use the group's child coordinate space (chOff/chExt), mapped onto its real box."""
    xfrm = grp._element.grpSpPr.find(qn("a:xfrm"))
    off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
    choff, chext = xfrm.find(qn("a:chOff")), xfrm.find(qn("a:chExt"))
    if None in (off, ext, choff, chext):
        return outer
    ox, oy, ew, eh = int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy"))
    cx, cy, cw, ch = int(choff.get("x")), int(choff.get("y")), int(chext.get("cx")) or 1, int(chext.get("cy")) or 1
    sx, sy = ew / cw, eh / ch

    def tf(x, y, w, h):
        return outer(ox + (x - cx) * sx, oy + (y - cy) * sy, w * sx, h * sy)
    return tf


def _para_size(p, scale):
    for path in ("a:r/a:rPr", "a:pPr/a:defRPr", "a:endParaRPr"):
        for el in p.findall("/".join(qn(t) for t in path.split("/"))):
            if el.get("sz"):
                return int(el.get("sz")) / 100 * scale
    return None


def _text_item(sh, is_title, is_body):
    body = sh.text_frame._txBody
    autofit = body.find(qn("a:bodyPr") + "/" + qn("a:normAutofit"))
    scale = int(autofit.get("fontScale")) / 100000 if autofit is not None and autofit.get("fontScale") else 1
    anchor = body.find(qn("a:bodyPr")).get("anchor") if body.find(qn("a:bodyPr")) is not None else None
    paras = []
    for p in body.findall(qn("a:p")):
        text = xml_text(p)
        ppr = p.find(qn("a:pPr"))
        bullet = False
        if text:
            if ppr is not None and ppr.find(qn("a:buNone")) is not None:
                bullet = False
            elif ppr is not None and (ppr.find(qn("a:buChar")) is not None or ppr.find(qn("a:buAutoNum")) is not None):
                bullet = True
            else:
                bullet = is_body
        rpr = p.find(qn("a:r") + "/" + qn("a:rPr"))
        paras.append({
            "text": text, "size": _para_size(p, scale) or (32 * scale if is_title else 18 * scale),
            "bold": is_title or (rpr is not None and rpr.get("b") in ("1", "true")),
            "level": int(ppr.get("lvl", 0)) if ppr is not None else 0,
            "align": {"ctr": "center", "r": "right", "just": "justify"}.get(ppr.get("algn") if ppr is not None else None, "left"),
            "bullet": bullet,
        })
    while paras and not paras[-1]["text"]:
        paras.pop()
    return {"type": "text", "paras": paras, "anchor": {"ctr": "middle", "b": "bottom"}.get(anchor, "top")} if paras else None


def _slide_items(shapes, tf, did, out):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            _slide_items(sh.shapes, _group_transform(sh, tf), did, out)
            continue
        if sh.left is None or sh.width is None:
            continue
        x, y, w, h = tf(sh.left, sh.top, sh.width, sh.height)
        box = {"x": x / EMU_PT, "y": y / EMU_PT, "w": w / EMU_PT, "h": h / EMU_PT, "rot": getattr(sh, "rotation", 0) or 0}
        item = None
        if _local(sh._element.tag) == "pic" and hasattr(sh, "image"):
            try:
                src = _picture(did, sh.image.blob, sh.image.ext)
            except Exception:
                src = None
            item = {"type": "image", "src": src,
                    "crop": [sh.crop_left, sh.crop_top, sh.crop_right, sh.crop_bottom] if src else None}
        elif getattr(sh, "has_table", False) and sh.has_table:
            rows = [[xml_text(c) for c in r.iter(qn("a:tc"))] for r in sh._element.iter(qn("a:tr"))]
            item = {"type": "table", "rows": rows}
        elif sh.has_text_frame and sh.text_frame.text.strip():
            ph = sh.placeholder_format.type if sh.is_placeholder else None
            is_title = ph in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
            is_body = ph in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT)
            item = _text_item(sh, is_title, is_body)
        if item:
            out.append({**box, **item})


def pptx_layout(path: Path, did: int) -> dict:
    prs = Presentation(path)
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        items = []
        _slide_items(slide.shapes, _identity, did, items)
        slides.append({"idx": i, "items": items})
    return {"kind": "slides", "width": prs.slide_width / EMU_PT, "height": prs.slide_height / EMU_PT, "slides": slides}


# ---------------------------------------------------------------- DOCX

def docx_layout(path: Path, did: int) -> dict:
    """Blocks in document order, each tagged with the unit (section) it belongs to, matching parsers.parse_docx."""
    blocks = docx_blocks(DocxDocument(path))
    groups, _ = docx_groups(blocks)
    unit_of = {n: gi for gi, g in enumerate(groups, 1) for n in range(g["first"], g["last"] + 1)}
    out = []
    for n, (kind, text, blobs, _eq, el) in enumerate(blocks, 1):
        b = {"unit": unit_of.get(n), "type": kind}
        if kind == "table":
            b["rows"] = [[xml_text(c) for c in r.iter("{*}tc")] for r in el.iter("{*}tr")]
        else:
            b["text"] = text
            style = el.find(f"{W_NS}pPr/{W_NS}pStyle")
            name = (style.get(f"{W_NS}val") or "").lower() if style is not None else ""
            if kind == "heading":
                b["level"] = 1 if name == "title" else int(name[7:]) + 1 if name.startswith("heading") and name[7:].isdigit() else 3
            if el.find(f"{W_NS}pPr/{W_NS}numPr") is not None:
                lvl = el.find(f"{W_NS}pPr/{W_NS}numPr/{W_NS}ilvl")
                b["list"] = int(lvl.get(f"{W_NS}val", 0)) if lvl is not None else 0
            jc = el.find(f"{W_NS}pPr/{W_NS}jc")
            if jc is not None and jc.get(f"{W_NS}val") in ("center", "right"):
                b["align"] = jc.get(f"{W_NS}val")
            b["images"] = [src for blob, ext in blobs if (src := _picture(did, blob, ext))]
            if not text and not b["images"]:
                continue
        out.append(b)
    return {"kind": "docx", "blocks": out}
