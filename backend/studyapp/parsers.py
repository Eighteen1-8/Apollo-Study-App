"""Native-Python document parsing into normalised units (slide / page / section).

Every parser returns a list of dicts:
    {idx, kind, label, text, equations, images: [asset paths], image_hashes: [sha256], reasons: [...], has_visual}
`reasons` is the automatic vision heuristic; the vision mode (off/auto/all) is applied later.
Rendering images for vision happens lazily in `render_for_vision`, straight from the original file.
"""
import hashlib
import io
import math
import re
from collections import Counter
from pathlib import Path

import pymupdf
from docx import Document as DocxDocument
from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .omml import omml_to_latex

PARSER_VERSION = "2"  # 2: SmartArt and chart text
M_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
R_EMBED = R_NS + "embed"
TEXT_EXTS = {".txt", ".md", ".csv", ".tsv", ".json", ".html", ".htm", ".xml", ".rtf", ".tex"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif", ".tif", ".tiff"}
SUPPORTED = {".pdf", ".pptx", ".docx"} | TEXT_EXTS | IMAGE_EXTS


def approx_tokens(text: str) -> int:
    return math.ceil(len(text) / 3.8)


def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _local(tag):
    return tag.rsplit("}", 1)[-1] if isinstance(tag, str) else ""


def xml_text(el) -> str:
    """Text of an OOXML element in document order. Equations become $LaTeX$; mc:Fallback is skipped
    (it holds a duplicate picture/text of the mc:Choice content)."""
    out = []

    def walk(e):
        name = _local(e.tag)
        if name == "Fallback":
            return
        if e.tag == M_NS + "oMath":
            out.append(" $" + omml_to_latex(e) + "$ ")
            return
        if name == "t" and e.text and not e.tag.startswith(M_NS):
            out.append(e.text)
        elif name in ("br", "tab"):
            out.append(" ")
        for c in e:
            walk(c)

    walk(el)
    return re.sub(r"[ \t]+", " ", "".join(out)).strip()


def _related_xml(part, rid):
    try:
        return etree.fromstring(part.rels[rid].target_part.blob)
    except Exception:
        return None


def graphic_text(el, part) -> list[str]:
    """Text that lives outside the slide/page XML, in related parts: SmartArt (diagram data) and charts."""
    out = []
    for gd in el.iter("{*}graphicData"):
        if any(_local(a.tag) == "Fallback" for a in gd.iterancestors()):
            continue
        uri = gd.get("uri", "")
        if uri.endswith("/diagram"):
            rel = gd.find("{*}relIds")
            x = _related_xml(part, rel.get(R_NS + "dm")) if rel is not None else None
            if x is None:
                continue
            items = []
            for pt in x.iter("{*}pt"):
                if pt.get("type") in ("parTrans", "sibTrans", "pres"):
                    continue
                t = pt.find("{*}t")
                line = " ".join(filter(None, (xml_text(p) for p in t.iter("{*}p")))) if t is not None else ""
                if line:
                    items.append(line)
            if items:
                out.append("[diagram]\n" + "\n".join(f"- {i}" for i in items))
        elif uri.endswith("/chart"):
            ch = gd.find("{*}chart")
            x = _related_xml(part, ch.get(R_NS + "id")) if ch is not None else None
            if x is None:
                continue
            titles = [" ".join(t.text for t in title.iter("{*}t") if t.text) for title in x.iter("{*}title")]
            series = [v.text for ser in x.iter("{*}ser") for tx in ser.findall("{*}tx") for v in tx.iter("{*}v") if v.text]
            cats = list(dict.fromkeys(v.text for cat in x.iter("{*}cat") for v in cat.iter("{*}v") if v.text))
            parts = [f"title: {'; '.join(t for t in titles if t)}" if any(titles) else "",
                     f"series: {', '.join(series)}" if series else "", f"categories: {', '.join(cats[:40])}" if cats else ""]
            if any(parts):
                out.append("[chart] " + " | ".join(p for p in parts if p))
    return out


def _count_eq(el):
    return sum(1 for m in el.iter(M_NS + "oMath") if not any(_local(a.tag) == "Fallback" for a in m.iterancestors()))


def _save_asset(asset_dir: Path, name: str, blob: bytes) -> str:
    asset_dir.mkdir(parents=True, exist_ok=True)
    (asset_dir / name).write_bytes(blob)
    return f"{asset_dir.name}/{name}"


# ---------------------------------------------------------------- PPTX

def _iter_shapes(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(sh.shapes)
        else:
            yield sh


def _pptx_pictures(slide):
    for sh in _iter_shapes(slide.shapes):
        # Filled picture placeholders are <p:pic> too but report shape_type PLACEHOLDER.
        if _local(sh._element.tag) == "pic" and hasattr(sh, "image"):
            try:
                yield sh, sh.image.blob, sh.image.ext
            except Exception:
                continue


def parse_pptx(path: Path, asset_dir: Path, s: dict):
    prs = Presentation(path)
    area = prs.slide_width * prs.slide_height
    raw = []
    for i, slide in enumerate(prs.slides, 1):
        tree = slide.shapes._spTree
        texts = []
        tables = list(tree.iter("{*}tbl"))
        for tbl in tables:
            rows = [" | ".join(xml_text(c) for c in r.iter("{*}tc")) for r in tbl.iter("{*}tr")]
            texts.append("[table]\n" + "\n".join(rows))
        for para in tree.iter("{*}p"):
            if _local(para.tag) != "p":
                continue
            if any(a in tables or _local(a.tag) == "Fallback" for a in para.iterancestors()):
                continue
            t = xml_text(para)
            if t:
                texts.append(t)
        texts += graphic_text(tree, slide.part)
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        pics = [{"hash": sha(blob), "ext": ext, "blob": blob, "area": (sh.width or 0) * (sh.height or 0) / area}
                for sh, blob, ext in _pptx_pictures(slide)]
        title = slide.shapes.title.text_frame.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
        text = "\n".join(texts) + (f"\n[speaker notes]\n{notes}" if notes else "")
        raw.append({"idx": i, "title": title, "text": text.strip(), "pics": pics, "equations": _count_eq(tree)})

    # Pictures repeated on most slides are template decoration (logos, banners): exact byte hash.
    seen = Counter(h for r in raw for h in {p["hash"] for p in r["pics"]})
    template = {h for h, n in seen.items() if len(raw) > 2 and n / len(raw) > 0.5}

    units = []
    for r in raw:
        content = [p for p in r["pics"] if p["hash"] not in template and p["area"] >= 0.02]
        words = len(r["text"].split())
        reasons = []
        if len(content) >= s["vision_min_images"]:
            reasons.append(f"{len(content)} images")
        if any(p["area"] > s["vision_big_image_area"] for p in content):
            reasons.append("large image")
        if content and words < s["vision_low_text_words"]:
            reasons.append(f"little text beside image ({words} words)")
        images = [_save_asset(asset_dir, f"s{r['idx']}_{k}.{p['ext']}", p["blob"]) for k, p in enumerate(content)]
        units.append({
            "idx": r["idx"], "kind": "slide",
            "label": f"Slide {r['idx']}" + (f" · {r['title'][:60]}" if r["title"] else ""),
            "text": r["text"], "equations": r["equations"], "images": images,
            "image_hashes": [p["hash"] for p in content], "reasons": reasons, "has_visual": bool(content),
        })
    return units, {}


def _font(size):
    try:
        return ImageFont.truetype("arial.ttf", size)
    except OSError:
        return ImageFont.load_default()


def compose_slide(path: Path, idx: int, width=1600) -> bytes:
    """Approximate slide render without PowerPoint/LibreOffice: pictures at their true position and size,
    text boxes drawn in place. Connectors/shape-drawn arrows are not reproduced."""
    prs = Presentation(path)
    slide = prs.slides[idx - 1]
    scale = width / prs.slide_width
    canvas = Image.new("RGB", (width, int(prs.slide_height * scale)), "white")
    draw = ImageDraw.Draw(canvas)
    for sh in _iter_shapes(slide.shapes):
        if sh.left is None or sh.width is None:
            continue
        box = (int(sh.left * scale), int(sh.top * scale), max(int(sh.width * scale), 1), max(int(sh.height * scale), 1))
        if _local(sh._element.tag) == "pic" and hasattr(sh, "image"):
            try:
                im = Image.open(io.BytesIO(sh.image.blob))
                im.load()
                canvas.paste(im.convert("RGB").resize(box[2:]), box[:2])
            except Exception:
                draw.rectangle([box[0], box[1], box[0] + box[2], box[1] + box[3]], outline="gray")
                draw.text((box[0] + 6, box[1] + 6), "[unrenderable image]", fill="gray", font=_font(18))
        elif sh.has_text_frame and sh.text_frame.text.strip():
            draw.multiline_text((box[0] + 4, box[1] + 4), sh.text_frame.text, fill="black", font=_font(22))
    buf = io.BytesIO()
    canvas.save(buf, "PNG")
    return buf.getvalue()


# ---------------------------------------------------------------- DOCX

def _docx_is_heading(p_el, text):
    style = p_el.find(".//{*}pStyle")
    name = (style.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val") or "").lower() if style is not None else ""
    if name.startswith("heading") or name == "title":
        return True
    if p_el.find(".//{*}outlineLvl") is not None:
        return True
    words = len(text.split())
    if not text or words > 12 or text.endswith((".", ",", ";")):
        return False
    runs = [r for r in p_el.iter("{*}r") if "".join(t.text or "" for t in r.iter("{*}t")).strip()]
    if not runs:
        return False
    bold = all(r.find("{*}rPr/{*}b") is not None and r.find("{*}rPr/{*}b").get(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val", "true") not in ("0", "false") for r in runs)
    sizes = [int(sz.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val", 0))
             for r in runs for sz in r.findall("{*}rPr/{*}sz")]
    big = bool(sizes) and min(sizes) >= 28  # half-points: 28 = 14pt
    return bold or big


def docx_blocks(doc) -> list[tuple]:
    """Body paragraphs and tables in order: (kind, text, image_blobs, equations, element)."""
    parts = doc.part.related_parts
    paras = []
    for el in doc.element.body.iterchildren():
        tag = _local(el.tag)
        if tag == "p":
            text = xml_text(el)
            blobs = []
            for blip in el.iter("{*}blip"):
                if any(_local(a.tag) == "Fallback" for a in blip.iterancestors()):
                    continue
                part = parts.get(blip.get(R_EMBED))
                if part is not None:
                    blobs.append((part.blob, Path(part.partname).suffix.lstrip(".") or "png"))
            kind = "heading" if _docx_is_heading(el, text) else "para"
            extra = graphic_text(el, doc.part)
            if extra:
                text = "\n".join([text, *extra]).strip()
            paras.append((kind, text, blobs, _count_eq(el), el))
        elif tag == "tbl":
            rows = [" | ".join(xml_text(c) for c in r.iter("{*}tc")) for r in el.iter("{*}tr")]
            paras.append(("table", "[table]\n" + "\n".join(rows), [], _count_eq(el), el))
    return paras


def docx_groups(paras) -> tuple[list[dict], bool]:
    """Group blocks into sections (one unit each): by heading, or ~350-token runs when there are no headings."""
    has_headings = any(p[0] == "heading" for p in paras)
    groups, cur = [], None
    for n, (kind, text, blobs, eq, _el) in enumerate(paras, 1):
        start_new = cur is None or (kind == "heading" if has_headings else
                                    approx_tokens("\n".join(cur["texts"])) > 350)
        if start_new:
            cur = {"heading": text if kind == "heading" else "", "first": n, "last": n, "texts": [], "blobs": [], "eq": 0}
            groups.append(cur)
        cur["last"] = n
        if text:
            cur["texts"].append(text)
        cur["blobs"] += blobs
        cur["eq"] += eq
    return groups, has_headings


def parse_docx(path: Path, asset_dir: Path, s: dict):
    groups, has_headings = docx_groups(docx_blocks(DocxDocument(path)))
    units = []
    for i, g in enumerate(groups, 1):
        text = "\n".join(g["texts"]).strip()
        if not text and not g["blobs"]:
            continue
        images = [_save_asset(asset_dir, f"d{i}_{k}.{ext}", blob) for k, (blob, ext) in enumerate(g["blobs"])]
        label = g["heading"][:70] if g["heading"] else f"Paragraphs {g['first']}–{g['last']}"
        units.append({
            "idx": i, "kind": "section", "label": label, "text": text, "equations": g["eq"],
            "images": images, "image_hashes": [sha(b) for b, _ in g["blobs"]],
            "reasons": [f"{len(images)} embedded image(s)"] if images else [], "has_visual": bool(images),
        })
    return units, {"headings_detected": has_headings}


# ---------------------------------------------------------------- PDF

def parse_pdf(path: Path, asset_dir: Path, s: dict):
    pdf = pymupdf.open(path)
    pages = []
    for p in pdf:
        text = p.get_text("text", sort=True)
        area = p.rect.width * p.rect.height
        rects = [r for info in p.get_images(full=True) for r in p.get_image_rects(info[0])]
        pages.append({"p": p, "text": text, "area": area, "rects": rects, "drawings": len(p.get_drawings())})

    # Deterministic header/footer removal: a line (surrounding whitespace stripped) that appears
    # character-for-character on >= 50% of pages is dropped.
    stripped, repeated = 0, set()
    if s["strip_repeated_pdf_lines"] and len(pages) >= 4:
        counts = Counter(l for pg in pages for l in {ln.strip() for ln in pg["text"].splitlines()} if len(l) >= 3)
        repeated = {l for l, n in counts.items() if n / len(pages) >= 0.5}

    units = []
    for pg in pages:
        lines = [ln.rstrip() for ln in pg["text"].splitlines()]
        kept = [ln for ln in lines if ln.strip() not in repeated]
        stripped += len(lines) - len(kept)
        text = re.sub(r"\n{3,}", "\n\n", "\n".join(kept)).strip()
        words = len(text.split())
        full = [r for r in pg["rects"] if r.width * r.height > 0.9 * pg["area"]]
        fig_area = sum(r.width * r.height for r in pg["rects"] if r not in full) / pg["area"]
        reasons = []
        if words < 20 and full:
            reasons.append("scanned page")
        elif not text:
            reasons.append("no text layer")
        if fig_area >= s["vision_pdf_figure_area"]:
            reasons.append(f"figure ({fig_area:.0%} of page)")
        if pg["drawings"] > s["vision_pdf_vector_paths"]:
            reasons.append(f"vector drawing ({pg['drawings']} paths)")
        n = pg["p"].number + 1
        units.append({
            "idx": n, "kind": "page", "label": f"Page {n}", "text": text, "equations": 0, "images": [],
            "image_hashes": [], "reasons": reasons, "has_visual": True,
        })
    return units, {"stripped_lines": stripped}


def render_pdf_page(path: Path, idx: int, dpi=150) -> bytes:
    return pymupdf.open(path)[idx - 1].get_pixmap(dpi=dpi).tobytes("png")


# ---------------------------------------------------------------- plain text & images

def parse_text(path: Path, asset_dir: Path, s: dict):
    raw = path.read_bytes().decode("utf-8", errors="replace")
    if path.suffix.lower() in (".html", ".htm", ".xml"):
        raw = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", raw, flags=re.S | re.I)
        raw = re.sub(r"<[^>]+>", " ", raw)
    blocks, cur = [], []
    for para in re.split(r"\n\s*\n", raw):
        cur.append(para.strip())
        if approx_tokens("\n\n".join(cur)) > 400:
            blocks.append("\n\n".join(cur))
            cur = []
    if cur:
        blocks.append("\n\n".join(cur))
    return [{"idx": i, "kind": "section", "label": f"Part {i}", "text": b.strip(), "equations": 0, "images": [],
             "image_hashes": [], "reasons": [], "has_visual": False} for i, b in enumerate(blocks, 1) if b.strip()], {}


def parse_image(path: Path, asset_dir: Path, s: dict):
    blob = path.read_bytes()
    img = _save_asset(asset_dir, f"img{path.suffix.lower()}", blob)
    return [{"idx": 1, "kind": "image", "label": "Image", "text": "", "equations": 0, "images": [img],
             "image_hashes": [sha(blob)], "reasons": ["standalone image"], "has_visual": True}], {}


def parse(path: Path, asset_dir: Path, settings: dict):
    ext = path.suffix.lower()
    if ext == ".pptx":
        return parse_pptx(path, asset_dir, settings)
    if ext == ".docx":
        return parse_docx(path, asset_dir, settings)
    if ext == ".pdf":
        return parse_pdf(path, asset_dir, settings)
    if ext in TEXT_EXTS:
        return parse_text(path, asset_dir, settings)
    if ext in IMAGE_EXTS:
        return parse_image(path, asset_dir, settings)
    raise ValueError(f"Unsupported file type {ext} (legacy .ppt/.doc: re-save as .pptx/.docx)")


def render_for_vision(path: Path, unit: dict, assets_root: Path) -> list[bytes]:
    """Images to send to the vision model for this unit."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        return [render_pdf_page(path, unit["idx"])]
    if ext == ".pptx":
        return [compose_slide(path, unit["idx"])]
    return [(assets_root / rel).read_bytes() for rel in unit["images"]]
